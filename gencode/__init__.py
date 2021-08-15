from jinja2 import Environment, FileSystemLoader
import jinja2_stringcase
import stringcase
import csv
import openpyxl
from gencode.internal import default
import gencode.text_helpers as text_helpers
import os

# Config Loading


class SourceConfig:
    def __init__(self) -> None:
        self.tables = []
        self.fields = []
        self.output_path = "output"
        self.template_path = "templates"

    def build(self) -> None:
        for table in self.tables:
            for field in self.fields:
                if field.table == table.name:
                    field.table_meta = table
                    table.fields.append(field)
        self.load_templates()

    def load_templates(self) -> None:
        file_loader = FileSystemLoader(self.template_path)
        env = Environment(loader=file_loader, extensions=[
            'jinja2_stringcase.StringCaseExtension'])
        env.globals["list_fmt"] = text_helpers.list_format
        self.jinja2 = env

    def get_template(self, name: str):
        """get a jinja2 template object linked to a file template."""
        return self.jinja2.get_template(name)

    def get_fields(self, table_name):
        """return a list of fields of specific table"""
        newlist = []
        for field in self.fields:
            if field.table == table_name:
                newlist.append(field)
            else:
                newlist
        return newlist


class Table:
    def __init__(self, name: str, description: str) -> None:
        self.name = name
        self.description = description
        self.fields = []


class Field:
    def __init__(self, table: str, name: str, dbtype: str, size: str, allownull: str, key: str, elm_type: str) -> None:
        self.table = table
        self.name = name
        self.dbtype = dbtype
        self.size = size
        self.key = key
        self.allownull = allownull
        self.elm_type = elm_type
        self.elm_type_func = stringcase.camelcase(elm_type)
        self.elm_name = stringcase.camelcase(name)
        self.elm_pascal = stringcase.pascalcase(name)
        self.elm_title = stringcase.titlecase(name)
        self.table_pascal_case = stringcase.pascalcase(table)


def load_from_xlsx(filename: str) -> SourceConfig:
    src = SourceConfig()
    wb = openpyxl.load_workbook(filename=filename)
    ws = wb['tables']
    header = True
    for row in ws.values:
        if header:
            header = False
            continue
        t1 = Table(name=default(row, 0),
                   description=default(row, 1))
        t1.row = row
        src.tables.append(t1)

    ws = wb['fields']
    header = True
    for row in ws.values:
        if header:
            header = False
            continue
        f1 = Field(
            table=default(row, 0),
            name=default(row, 1),
            dbtype=default(row, 2),
            size=default(row, 3),
            key=default(row, 4),
            allownull=default(row, 5),
            elm_type=default(row, 6))
        f1.row = row
        src.fields.append(f1)

    src.build()
    return src


def load_from_csv(path: str) -> SourceConfig:
    src = SourceConfig()
    # Load Table config
    with open(path+"/tables.csv", mode='r', encoding='utf-8-sig') as csv_file:
        csv_reader = csv.DictReader(csv_file)
        for row in csv_reader:
            t = Table(name=row['name'], description=row['description'])
            t.row = row
            src.tables.append(t)

    # Load Field config
    with open(path+"/fields.csv", mode='r', encoding='utf-8-sig') as csv_file:
        csv_reader = csv.DictReader(csv_file)
        for row in csv_reader:
            f1 = Field(
                table=row['table'],
                name=row['name'],
                dbtype=row['dbtype'],
                size=row['size'],
                key=row['key'],
                allownull=row['allownull'],
                elm_type=row['elm_type'])
            f1.row = row
            src.fields.append(f1)
    src.build()
    return src


# helper


def write_file(filename, contents):
    filename = os.path.realpath(filename)
    directory = os.path.dirname(filename)
    if not os.path.exists(directory):
        os.makedirs(directory)
    f = open(filename, "w")
    f.write(contents)
    f.close()
